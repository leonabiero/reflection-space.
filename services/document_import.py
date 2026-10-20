"""
Document import
=================

Lets a professional bring an existing document into the Documentation page
instead of retyping it: a plain text file (.txt), a Word file (.docx) or a
PDF that contains real text.

What it does -- and deliberately does not do
---------------------------------------------
- It only EXTRACTS TEXT. The uploaded file itself is never stored anywhere;
  only the text the professional reviews and chooses to save is stored, the
  same as if they had typed it. From there the normal path applies (the text
  is anonymised before any AI sees it).
- It never guesses: a scanned PDF (a picture of paper, no text layer), a
  password-protected file, or a file that is too large or too long is
  refused with a clear reason, not half-read.
- It never logs file names or contents.

Libraries (approved by the product owner, October 2026)
--------------------------------------------------------
  pypdf        reads text from PDF files
  python-docx  reads text from Word (.docx) files
Both are free, open source and run inside the app (no outside service, no
cost). They are imported only when a PDF / Word file is actually opened, so
the rest of the app is unaffected if they are missing.

Limits (all in one place so they can be reviewed)
---------------------------------------------------
"""

import io
import re
import unicodedata
import zipfile

MAX_FILE_BYTES = 5 * 1024 * 1024          # 5 MB uploaded file
MAX_UNZIPPED_BYTES = 40 * 1024 * 1024     # guards against "zip bombs" in .docx
MAX_PDF_PAGES = 60
MAX_CHARS = 30000                          # about 10-15 printed pages of text

ALLOWED_EXTENSIONS = ("txt", "pdf", "docx")

# Reasons a file can be refused. The page maps each to a translated message.
UNSUPPORTED = "unsupported"
TOO_LARGE = "too_large"
TOO_LONG = "too_long"
TOO_MANY_PAGES = "too_many_pages"
ENCRYPTED = "encrypted"
NO_TEXT = "no_text"            # empty file, or a scan with no text layer
UNREADABLE = "unreadable"
UNAVAILABLE = "unavailable"    # the reading library is not installed


class DocumentImportError(Exception):
    """A file could not be turned into text. `reason` is one of the codes above."""

    def __init__(self, reason):
        super().__init__(reason)
        self.reason = reason


def file_extension(filename):
    name = (filename or "").rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


def clean_text(text):
    """Make extracted text tidy and safe to show: normal Unicode form, no
    control characters or zero-width marks, Windows line endings made plain,
    and no more than one blank line in a row."""
    text = unicodedata.normalize("NFC", text or "")
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = "".join(
        ch for ch in text
        if ch in "\n\t" or (unicodedata.category(ch) not in ("Cc", "Cf", "Cs", "Co"))
    )
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _from_txt(data):
    # UTF-16 decodes almost any bytes without error, so it is only tried
    # when the file announces itself with the UTF-16 byte-order mark.
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        try:
            return data.decode("utf-16")
        except UnicodeError:
            raise DocumentImportError(UNREADABLE)
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeError:
            continue
    raise DocumentImportError(UNREADABLE)


def _from_docx(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_UNZIPPED_BYTES:
                raise DocumentImportError(TOO_LARGE)
    except zipfile.BadZipFile:
        raise DocumentImportError(UNREADABLE)

    try:
        import docx
    except Exception:
        raise DocumentImportError(UNAVAILABLE)

    try:
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    value = cell.text.strip()
                    if value and (not cells or cells[-1] != value):   # merged cells repeat
                        cells.append(value)
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except DocumentImportError:
        raise
    except Exception:
        raise DocumentImportError(UNREADABLE)


def _from_pdf(data):
    try:
        import pypdf
    except Exception:
        raise DocumentImportError(UNAVAILABLE)

    try:
        reader = pypdf.PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            # An "empty password" PDF opens; a real password does not.
            try:
                opened = reader.decrypt("")
            except Exception:
                opened = 0
            if not opened:
                raise DocumentImportError(ENCRYPTED)
        if len(reader.pages) > MAX_PDF_PAGES:
            raise DocumentImportError(TOO_MANY_PAGES)
        parts = []
        total = 0
        for page in reader.pages:
            page_text = page.extract_text() or ""
            parts.append(page_text)
            total += len(page_text)
            if total > MAX_CHARS * 2:      # stop early; it will be refused as too long
                break
        return "\n\n".join(parts)
    except DocumentImportError:
        raise
    except Exception:
        raise DocumentImportError(UNREADABLE)


def extract_text(filename, data):
    """
    Turn an uploaded file into clean text.

    `filename`: only its ending (.txt / .pdf / .docx) is used.
    `data`: the file's bytes.

    Returns the text. Raises DocumentImportError(reason) -- never returns
    partial or invented text.
    """
    extension = file_extension(filename)
    if extension not in ALLOWED_EXTENSIONS:
        raise DocumentImportError(UNSUPPORTED)
    if not isinstance(data, (bytes, bytearray)) or len(data) == 0:
        raise DocumentImportError(NO_TEXT)
    if len(data) > MAX_FILE_BYTES:
        raise DocumentImportError(TOO_LARGE)

    data = bytes(data)
    if extension == "txt":
        raw = _from_txt(data)
    elif extension == "docx":
        raw = _from_docx(data)
    else:
        raw = _from_pdf(data)

    text = clean_text(raw)
    if not text:
        raise DocumentImportError(NO_TEXT)
    if len(text) > MAX_CHARS:
        raise DocumentImportError(TOO_LONG)
    return text
