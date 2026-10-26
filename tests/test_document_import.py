"""Tests for reading uploaded documents (txt / Word / PDF)."""
import io
import re
import unittest
import zipfile
from pathlib import Path

from services import document_import as di

ROOT = Path(__file__).resolve().parents[1]


def _pdf_with_text(text):
    """A tiny hand-made PDF containing `text` (ASCII) on one page."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objs) + 1, xref)
    return out


def _docx_bytes(paragraphs, table=None):
    import docx
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.rows[r].cells[c].text = value
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def reason(filename, data):
    try:
        di.extract_text(filename, data)
    except di.DocumentImportError as exc:
        return exc.reason
    return None


class TextFileTests(unittest.TestCase):
    def test_utf8_cp1252_and_utf16(self):
        self.assertEqual(di.extract_text("a.txt", "Árbol ñandú".encode("utf-8")), "Árbol ñandú")
        self.assertEqual(di.extract_text("a.TXT", "Árbol ñandú".encode("cp1252")), "Árbol ñandú")
        self.assertEqual(di.extract_text("a.txt", "Zuri «ñ»".encode("utf-16")), "Zuri «ñ»")

    def test_tidying(self):
        raw = "Uno\r\n\r\n\r\n\r\nDos\u200b  \x00fin".encode("utf-8")
        text = di.extract_text("a.txt", raw)
        self.assertNotIn("\r", text)
        self.assertNotIn("\n\n\n", text)
        self.assertNotIn("\u200b", text)
        self.assertNotIn("\x00", text)


class WordTests(unittest.TestCase):
    def test_paragraphs_and_tables(self):
        data = _docx_bytes(["Hola ñ.", "Segunda."], table=[["A", "B"], ["C", "C"]])
        text = di.extract_text("n.docx", data)
        self.assertIn("Hola ñ.", text)
        self.assertIn("Segunda.", text)
        self.assertIn("A | B", text)
        self.assertIn("\nC", text)

    def test_corrupt_and_zip_bomb_are_refused(self):
        self.assertEqual(reason("n.docx", b"not a zip at all"), di.UNREADABLE)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("word/document.xml", b"0" * (di.MAX_UNZIPPED_BYTES + 1024))
        self.assertLess(len(buf.getvalue()), di.MAX_FILE_BYTES)
        self.assertEqual(reason("n.docx", buf.getvalue()), di.TOO_LARGE)


class PdfTests(unittest.TestCase):
    def test_reads_real_text(self):
        self.assertIn("Visita domiciliaria", di.extract_text("n.pdf", _pdf_with_text("Visita domiciliaria")))

    def test_scanned_or_blank_pdf_is_refused_not_invented(self):
        import pypdf
        w = pypdf.PdfWriter(); w.add_blank_page(612, 792)
        buf = io.BytesIO(); w.write(buf)
        self.assertEqual(reason("scan.pdf", buf.getvalue()), di.NO_TEXT)

    def test_password_protected_pdf_is_refused(self):
        import pypdf
        w = pypdf.PdfWriter(); w.append(pypdf.PdfReader(io.BytesIO(_pdf_with_text("secreto"))))
        try:
            w.encrypt("pw")
        except Exception:
            self.skipTest("encryption support not available here")
        buf = io.BytesIO(); w.write(buf)
        self.assertEqual(reason("p.pdf", buf.getvalue()), di.ENCRYPTED)

    def test_too_many_pages(self):
        import pypdf
        w = pypdf.PdfWriter()
        for _ in range(di.MAX_PDF_PAGES + 1):
            w.add_blank_page(100, 100)
        buf = io.BytesIO(); w.write(buf)
        self.assertEqual(reason("big.pdf", buf.getvalue()), di.TOO_MANY_PAGES)

    def test_garbage_pdf(self):
        self.assertEqual(reason("x.pdf", b"%PDF-1.4 garbage"), di.UNREADABLE)


class LimitTests(unittest.TestCase):
    def test_unsupported_empty_large_long(self):
        self.assertEqual(reason("x.exe", b"abc"), di.UNSUPPORTED)
        self.assertEqual(reason("x.doc", b"abc"), di.UNSUPPORTED)
        self.assertEqual(reason("noextension", b"abc"), di.UNSUPPORTED)
        self.assertEqual(reason("x.txt", b""), di.NO_TEXT)
        self.assertEqual(reason("x.txt", b"   \n  "), di.NO_TEXT)
        self.assertEqual(reason("x.txt", b"a" * (di.MAX_FILE_BYTES + 1)), di.TOO_LARGE)
        self.assertEqual(reason("x.txt", ("palabra " * 6000).encode()), di.TOO_LONG)
        self.assertIsNone(reason("x.txt", b"a" * 1000))

    def test_path_in_filename_does_not_confuse_extension(self):
        self.assertEqual(di.file_extension("C:\\docs\\nota.final.DOCX"), "docx")


class WiringTests(unittest.TestCase):
    def test_requirements_pin_both_tools(self):
        req = (ROOT / "requirements.txt").read_text()
        self.assertRegex(req, r"(?m)^pypdf==\d")
        self.assertRegex(req, r"(?m)^python-docx==\d")

    def test_upload_limit_matches_server_setting(self):
        cfg = (ROOT / ".streamlit" / "config.toml").read_text()
        mb = int(re.search(r"maxUploadSize\s*=\s*(\d+)", cfg).group(1))
        self.assertEqual(mb * 1024 * 1024, di.MAX_FILE_BYTES)

    def test_translations_complete_and_page_uses_module(self):
        from services.locales.es import LANG_ES
        from services.locales.en import LANG_EN
        from services.locales.eu import LANG_EU
        reasons = {di.UNSUPPORTED, di.TOO_LARGE, di.TOO_LONG, di.TOO_MANY_PAGES, di.ENCRYPTED,
                   di.NO_TEXT, di.UNREADABLE, di.UNAVAILABLE}
        for lang in (LANG_ES, LANG_EN, LANG_EU):
            for key in ("doc_import_label", "doc_import_help", "doc_import_loaded"):
                self.assertTrue(lang[key])
            self.assertEqual(set(lang["doc_import_errors"]), reasons)
        # Since the reflection-only redesign (REFLECTIONS.md) the upload lives
        # on the single Reflection Space screen; the old Documentation page
        # is only a redirect.
        page = (ROOT / "pages" / "reflection_space.py").read_text(encoding="utf-8")
        self.assertIn("document_import.extract_text", page)

    def test_module_never_logs_or_stores(self):
        source = (ROOT / "services" / "document_import.py").read_text(encoding="utf-8")
        code = re.sub(r'""".*?"""', "", source, flags=re.S)
        self.assertNotIn("logger", code)
        self.assertNotIn("open(", code)


if __name__ == "__main__":
    unittest.main()
